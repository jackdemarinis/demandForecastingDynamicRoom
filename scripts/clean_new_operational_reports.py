from __future__ import annotations

import json
import math
import re
import subprocess
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = REPO_ROOT / "Data"
CLEANED_DIR = REPO_ROOT / "cleaned_data"
STAGING_DIR = Path(tempfile.gettempdir()) / "csc561_xls_staging"

RESERVATION_OUTPUT = CLEANED_DIR / "reservation_forecast_daily_cleaned.csv"
RESERVATION_DUPES = CLEANED_DIR / "reservation_forecast_duplicates_removed.csv"
RESERVATION_SUMMARY = CLEANED_DIR / "reservation_forecast_cleaning_summary.json"

BLOCK_DAILY_OUTPUT = CLEANED_DIR / "block_roomsales_daily_totals_cleaned.csv"
BLOCK_CLASS_OUTPUT = CLEANED_DIR / "block_roomsales_room_class_daily_cleaned.csv"
BLOCK_TYPE_OUTPUT = CLEANED_DIR / "block_roomsales_room_type_daily_cleaned.csv"
BLOCK_ADJUSTMENT_OUTPUT = CLEANED_DIR / "block_roomsales_adjustments_cleaned.csv"
BLOCK_DUPES = CLEANED_DIR / "block_roomsales_duplicates_removed.csv"
BLOCK_SUMMARY = CLEANED_DIR / "block_roomsales_cleaning_summary.json"

ROOM_TYPE_TO_CLASS = {
    "HKNS": "Standard",
    "KFNS": "Standard",
    "KNS": "Standard",
    "QQNS": "Standard",
    "DKSNS": "Suite",
    "HKOBNS": "Suite",
    "KESNS": "Suite",
    "KOBNS": "Suite",
    "KSSNS": "Suite",
    "QQOBNS": "Suite",
}


def collapse_spaces(value: object | None) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date_range(text: str) -> tuple[date, date]:
    start_text, end_text = [piece.strip() for piece in str(text).split("-")]
    return (
        datetime.strptime(start_text, "%m/%d/%Y").date(),
        datetime.strptime(end_text, "%m/%d/%Y").date(),
    )


def iso_date(value: object | None) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def as_int(value: object | None) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        return int(round(float(text)))
    return int(round(float(value)))


def as_number(value: object | None, digits: int | None = None) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, str):
        text = value.strip().replace("$", "").replace(",", "")
        if not text:
            return None
        number = float(text)
    else:
        number = float(value)
    return round(number, digits) if digits is not None else number


def pct_to_rate(value: object | None, *, allow_invalid: bool = False) -> float | None:
    if value in (None, ""):
        return None
    try:
        if isinstance(value, str):
            text = value.strip().replace("%", "").replace(",", "")
            if not text:
                return None
            number = float(text)
        else:
            number = float(value)
    except Exception:
        return None
    if abs(number) > 1:
        number /= 100.0
    if not allow_invalid and (number < 0 or number > 1):
        return None
    return round(number, 4)


def raw_xls_files() -> list[Path]:
    files = [
        path
        for path in DATA_DIR.glob("*.xls")
        if path.name.startswith("ReservationForecast") or path.name.startswith("Block")
    ]
    return sorted(files, key=lambda path: path.name.lower())


def ensure_staging_files(files: list[Path]) -> dict[str, Path]:
    STAGING_DIR.mkdir(parents=True, exist_ok=True)
    output_map: dict[str, Path] = {}
    items: list[tuple[Path, Path]] = []
    for raw_path in files:
        xlsx_path = STAGING_DIR / f"{raw_path.stem}.xlsx"
        items.append((raw_path, xlsx_path))
        output_map[raw_path.name] = xlsx_path
    script_lines = [
        "$ErrorActionPreference='Stop'",
        "$excel = New-Object -ComObject Excel.Application",
        "$excel.Visible = $false",
        "$excel.DisplayAlerts = $false",
        "try {",
    ]
    for raw_path, xlsx_path in items:
        raw_text = str(raw_path).replace("'", "''")
        xlsx_text = str(xlsx_path).replace("'", "''")
        script_lines.extend(
            [
                "  $wb = $null",
                "  try {",
                f"    $wb = $excel.Workbooks.Open('{raw_text}')",
                f"    $wb.SaveAs('{xlsx_text}', 51)",
                "  } finally {",
                "    if ($wb -ne $null) {",
                "      $wb.Close($false) | Out-Null",
                "      [System.Runtime.Interopservices.Marshal]::ReleaseComObject($wb) | Out-Null",
                "    }",
                "  }",
            ]
        )
    script_lines.extend(
        [
            "} finally {",
            "  $excel.Quit() | Out-Null",
            "  [System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null",
            "}",
        ]
    )
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as handle:
        handle.write("\n".join(script_lines))
        script_path = Path(handle.name)
    try:
        subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script_path),
            ],
            check=True,
            cwd=str(REPO_ROOT),
        )
    finally:
        script_path.unlink(missing_ok=True)
    return output_map


def missing_daily_dates(dates: list[str]) -> list[str]:
    parsed = sorted(pd.to_datetime(pd.Series(dates)).dt.date.tolist())
    if not parsed:
        return []
    present = set(parsed)
    current = parsed[0]
    missing: list[str] = []
    while current <= parsed[-1]:
        if current not in present:
            missing.append(current.isoformat())
        current += timedelta(days=1)
    return missing


def parse_reservation_forecast(staging_map: dict[str, Path]) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    records: list[dict] = []
    totals_rows_removed = 0
    file_meta: list[dict] = []

    for raw_name, xlsx_path in sorted(staging_map.items()):
        if not raw_name.startswith("ReservationForecast"):
            continue
        ws = load_workbook(xlsx_path, data_only=True, read_only=True).active
        period_text = next(
            cell for cell in next(ws.iter_rows(min_row=2, max_row=2, values_only=True)) if cell
        )
        period_start, period_end = parse_date_range(str(period_text))
        file_meta.append(
            {
                "source_file": raw_name,
                "source_sheet": ws.title,
                "report_period_start": period_start.isoformat(),
                "report_period_end": period_end.isoformat(),
                "used_rows": int(ws.max_row),
                "used_columns": int(ws.max_column),
            }
        )
        for row_idx, row in enumerate(ws.iter_rows(min_row=10, values_only=True), start=10):
            business_date = iso_date(row[0])
            text = collapse_spaces(row[0])
            if business_date:
                records.append(
                    {
                        "source_file": raw_name,
                        "source_sheet": ws.title,
                        "source_row": row_idx,
                        "report_period_start": period_start.isoformat(),
                        "report_period_end": period_end.isoformat(),
                        "business_date": business_date,
                        "arrivals": as_int(row[3]),
                        "departures": as_int(row[4]),
                        "stayovers": as_int(row[5]),
                        "projected_walk_ins": as_int(row[6]),
                        "projected_cxl_no_show_rate": pct_to_rate(row[7]),
                        "rooms_sold": as_int(row[8]),
                        "total_comp_rooms": as_int(row[9]),
                        "out_of_order_rooms": as_int(row[10]),
                        "occupancy_rate_reported": pct_to_rate(row[11]),
                        "room_revenue_net": as_number(row[12], 2),
                        "adr_reported": as_number(row[13], 4),
                        "revpar_reported": as_number(row[14], 4),
                        "transient_rooms": as_int(row[15]),
                        "group_rooms_picked_up": as_int(row[16]),
                        "group_rooms_not_picked_up": as_int(row[17]),
                        "group_room_revenue_net": as_number(row[18], 2),
                        "group_adr_reported": as_number(row[19], 4),
                        "group_non_gtd_arrivals": as_int(row[20]),
                        "group_gtd_arrivals": as_int(row[21]),
                        "group_comp_rooms": as_int(row[22]),
                        "total_group_guests_in_house": as_int(row[24]),
                    }
                )
            elif text == "Totals":
                totals_rows_removed += 1

    frame = pd.DataFrame(records).sort_values(["business_date", "source_file", "source_row"]).reset_index(
        drop=True
    )
    duplicate_mask = frame.duplicated(subset=["business_date"], keep="first")
    duplicate_rows = frame.loc[
        duplicate_mask, ["source_file", "source_sheet", "source_row", "business_date"]
    ].copy()
    if duplicate_rows.empty:
        duplicate_rows = pd.DataFrame(
            columns=[
                "dataset_name",
                "source_file",
                "source_sheet",
                "source_row",
                "business_date",
                "dedupe_key_columns",
                "dedupe_key",
                "kept_source_file",
                "kept_source_sheet",
                "kept_source_row",
                "drop_reason",
            ]
        )
    else:
        kept = frame.drop_duplicates(subset=["business_date"], keep="first").set_index("business_date")
        duplicate_rows["dataset_name"] = "reservation_forecast_daily"
        duplicate_rows["dedupe_key_columns"] = "business_date"
        duplicate_rows["dedupe_key"] = duplicate_rows["business_date"]
        duplicate_rows["kept_source_file"] = duplicate_rows["business_date"].map(kept["source_file"])
        duplicate_rows["kept_source_sheet"] = duplicate_rows["business_date"].map(kept["source_sheet"])
        duplicate_rows["kept_source_row"] = duplicate_rows["business_date"].map(kept["source_row"])
        duplicate_rows["drop_reason"] = "duplicate daily row with same business_date after cleaning"
        duplicate_rows = duplicate_rows[
            [
                "dataset_name",
                "source_file",
                "source_sheet",
                "source_row",
                "business_date",
                "dedupe_key_columns",
                "dedupe_key",
                "kept_source_file",
                "kept_source_sheet",
                "kept_source_row",
                "drop_reason",
            ]
        ]
    clean = frame.drop_duplicates(subset=["business_date"], keep="first").copy()

    clean.to_csv(RESERVATION_OUTPUT, index=False)
    duplicate_rows.to_csv(RESERVATION_DUPES, index=False)

    summary = {
        "dataset_family": "reservation_forecast",
        "raw_files_used": [meta["source_file"] for meta in file_meta],
        "cleaned_files_created": [
            RESERVATION_OUTPUT.name,
            RESERVATION_DUPES.name,
            RESERVATION_SUMMARY.name,
        ],
        "report_layout": {
            "sheet_name": "ReservationForecast",
            "header_rows_before_daily_data": 9,
            "daily_data_row_rule": "Rows beginning at row 10 where column A is a valid Excel date were treated as daily observations.",
            "trailing_totals_rule": "Rows where column A equals Totals were excluded from the cleaned daily table.",
        },
        "date_coverage": {
            "start": clean["business_date"].min(),
            "end": clean["business_date"].max(),
            "daily_rows_cleaned": int(len(clean)),
            "missing_business_dates_after_cleaning": missing_daily_dates(clean["business_date"].tolist()),
            "per_file": file_meta,
        },
        "transformations": [
            "Converted the report row date in column A from Excel date serial / datetime storage to ISO YYYY-MM-DD business_date.",
            "Dropped the two unnamed blank spacer columns and kept only named metric columns plus source audit columns.",
            "Normalized percent-like fields (Projected Cxl/No-Shows and Occ%) to decimal rates rounded to 4 decimals.",
            "Normalized monetary totals to numeric values rounded to 2 decimals and ADR / RevPAR style metrics to numeric values rounded to 4 decimals.",
            "Converted count fields to nullable integers.",
            "Removed the trailing Totals row from each workbook before deduplication.",
        ],
        "deduplication": {
            "logical_key": ["business_date"],
            "duplicate_rows_removed": int(len(duplicate_rows)),
            "note": "No duplicate business_date rows were present across the 2023 and 2024 ReservationForecast files after cleaning.",
        },
        "rows_removed_or_excluded": {
            "totals_rows_removed": totals_rows_removed,
            "duplicate_rows_removed": int(len(duplicate_rows)),
        },
        "data_quality_notes": [
            "Projected Walk-Ins and Projected Cxl/No-Shows are present as columns but every observed daily value in the 2023 and 2024 files is zero.",
            "ReservationForecast ADR and RevPAR are not explicitly labeled Gross or Net in the workbook; the cleaned file preserves the report labels as adr_reported and revpar_reported.",
        ],
        "new_data_capabilities_from_this_family": [
            "Daily arrivals, departures, and stayovers for 2023-2024.",
            "Daily rooms sold, out-of-order rooms, transient rooms, group pickup counts, room revenue, ADR, and RevPAR for 2023-2024.",
            "Daily group-specific revenue and arrival fields not present in the guest-list files.",
        ],
    }
    RESERVATION_SUMMARY.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return clean, duplicate_rows, summary


def parse_block_roomsales(staging_map: dict[str, Path]) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, dict]:
    daily_records: list[dict] = []
    class_records: list[dict] = []
    type_records: list[dict] = []
    adjustment_raw_records: list[dict] = []
    file_meta: list[dict] = []
    summary_section_rows_removed = 0
    footnote_rows_removed = 0

    for raw_name, xlsx_path in sorted(staging_map.items()):
        if not raw_name.startswith("Block"):
            continue
        ws = load_workbook(xlsx_path, data_only=True, read_only=True).active
        period_text = next(
            cell for cell in next(ws.iter_rows(min_row=2, max_row=2, values_only=True)) if cell
        )
        period_start, period_end = parse_date_range(str(period_text))
        total_rooms_text = next(
            cell for cell in next(ws.iter_rows(min_row=10, max_row=10, values_only=True)) if cell
        )
        period_total_rooms = int(re.search(r"(\d+)$", str(total_rooms_text)).group(1))
        report_day_count = (period_end - period_start).days + 1
        property_total_rooms = period_total_rooms // report_day_count
        file_meta.append(
            {
                "source_file": raw_name,
                "source_sheet": ws.title,
                "report_period_start": period_start.isoformat(),
                "report_period_end": period_end.isoformat(),
                "used_rows": int(ws.max_row),
                "used_columns": int(ws.max_column),
                "report_period_total_rooms": period_total_rooms,
                "property_total_rooms_inferred": property_total_rooms,
            }
        )

        current_date: str | None = None
        in_summary_section = False
        for row_idx, row in enumerate(ws.iter_rows(min_row=11, values_only=True), start=11):
            business_date = iso_date(row[0])
            text = collapse_spaces(row[0])
            payload = {
                "source_file": raw_name,
                "source_sheet": ws.title,
                "source_row": row_idx,
                "report_period_start": period_start.isoformat(),
                "report_period_end": period_end.isoformat(),
                "property_total_rooms": property_total_rooms,
            }
            if business_date:
                current_date = business_date
                daily_records.append(
                    {
                        **payload,
                        "business_date": business_date,
                        "occupancy_rate_reported": pct_to_rate(row[3]),
                        "rooms_sold": as_int(row[5]),
                        "transient_rooms": as_int(row[6]),
                        "group_rooms": as_int(row[8]),
                        "day_use_rooms": as_int(row[9]),
                        "comp_rooms": as_int(row[10]),
                        "out_of_order_rooms": as_int(row[11]),
                        "on_market_rooms": as_int(row[12]),
                        "room_revenue_gross": as_number(row[13], 2),
                        "adr_gross": as_number(row[14], 4),
                        "revpar_gross": as_number(row[15], 4),
                        "room_revenue_net": as_number(row[17], 2),
                        "adr_net": as_number(row[19], 4),
                        "revpar_net": as_number(row[21], 4),
                    }
                )
                continue
            if not text:
                continue
            if text.startswith("Summary - "):
                in_summary_section = True
                summary_section_rows_removed += 1
                continue
            if text.startswith("*Totals for past dates"):
                footnote_rows_removed += 1
                continue
            if in_summary_section:
                summary_section_rows_removed += 1
                continue
            if current_date is None:
                continue

            row_payload = {
                **payload,
                "business_date": current_date,
                "occupancy_rate_reported": pct_to_rate(row[3]),
                "rooms_sold": as_int(row[5]),
                "transient_rooms": as_int(row[6]),
                "group_rooms": as_int(row[8]),
                "day_use_rooms": as_int(row[9]),
                "comp_rooms": as_int(row[10]),
                "out_of_order_rooms": as_int(row[11]),
                "on_market_rooms": as_int(row[12]),
                "room_revenue_gross": as_number(row[13], 2),
                "adr_gross": as_number(row[14], 4),
                "revpar_gross": as_number(row[15], 4),
                "room_revenue_net": as_number(row[17], 2),
                "adr_net": as_number(row[19], 4),
                "revpar_net": as_number(row[21], 4),
            }
            if text in {"Standard", "Suite"}:
                class_records.append({**row_payload, "room_class": text})
            elif text == "***":
                adjustment_raw_records.append(
                    {
                        **row_payload,
                        "adjustment_label_raw": text,
                        "occupancy_rate_raw": row[3],
                    }
                )
            else:
                type_records.append(
                    {
                        **row_payload,
                        "room_class": ROOM_TYPE_TO_CLASS[text],
                        "room_type_code": text,
                    }
                )

    daily_df = pd.DataFrame(daily_records).sort_values(["business_date", "source_file", "source_row"]).reset_index(
        drop=True
    )
    class_df = pd.DataFrame(class_records).sort_values(
        ["business_date", "room_class", "source_file", "source_row"]
    ).reset_index(drop=True)
    type_df = pd.DataFrame(type_records).sort_values(
        ["business_date", "room_class", "room_type_code", "source_file", "source_row"]
    ).reset_index(drop=True)
    adjustment_raw_df = pd.DataFrame(adjustment_raw_records).sort_values(
        ["business_date", "source_file", "source_row"]
    ).reset_index(drop=True)

    duplicate_rows: list[dict] = []
    for dataset_name, frame, subset in [
        ("block_roomsales_daily_totals", daily_df, ["business_date"]),
        ("block_roomsales_room_class_daily", class_df, ["business_date", "room_class"]),
        ("block_roomsales_room_type_daily", type_df, ["business_date", "room_type_code"]),
    ]:
        duplicate_mask = frame.duplicated(subset=subset, keep="first")
        if duplicate_mask.any():
            kept = frame.drop_duplicates(subset=subset, keep="first").set_index(subset)
            for _, row in frame.loc[duplicate_mask].iterrows():
                key = tuple(row[column] for column in subset)
                kept_row = kept.loc[key]
                duplicate_rows.append(
                    {
                        "dataset_name": dataset_name,
                        "source_file": row["source_file"],
                        "source_sheet": row["source_sheet"],
                        "source_row": int(row["source_row"]),
                        "business_date": row["business_date"],
                        "dedupe_key_columns": "|".join(subset),
                        "dedupe_key": "|".join(str(value) for value in key),
                        "kept_source_file": kept_row["source_file"],
                        "kept_source_sheet": kept_row["source_sheet"],
                        "kept_source_row": int(kept_row["source_row"]),
                        "drop_reason": "duplicate cleaned block row with matching logical key",
                    }
                )
            frame.drop_duplicates(subset=subset, keep="first", inplace=True)

    adjustment_group_cols = [
        "business_date",
        "adjustment_label_raw",
        "rooms_sold",
        "transient_rooms",
        "group_rooms",
        "day_use_rooms",
        "comp_rooms",
        "out_of_order_rooms",
        "on_market_rooms",
        "room_revenue_gross",
        "adr_gross",
        "revpar_gross",
        "room_revenue_net",
        "adr_net",
        "revpar_net",
    ]
    adjustment_clean_records: list[dict] = []
    if not adjustment_raw_df.empty:
        adjustment_raw_df["occupancy_rate_raw_text"] = adjustment_raw_df["occupancy_rate_raw"].map(
            lambda value: "" if value is None else str(value)
        )
        grouped = adjustment_raw_df.groupby(adjustment_group_cols, dropna=False, sort=True)
        for _, group in grouped:
            group = group.sort_values(["source_file", "source_row"]).reset_index(drop=True)
            kept = group.iloc[0]
            adjustment_clean_records.append(
                {
                    "source_file": kept["source_file"],
                    "source_sheet": kept["source_sheet"],
                    "source_row": int(kept["source_row"]),
                    "report_period_start": kept["report_period_start"],
                    "report_period_end": kept["report_period_end"],
                    "property_total_rooms": int(kept["property_total_rooms"]),
                    "business_date": kept["business_date"],
                    "adjustment_label_raw": kept["adjustment_label_raw"],
                    "occupancy_rate_reported": None,
                    "occupancy_rate_raw_values_observed": "|".join(
                        sorted(
                            value
                            for value in set(group["occupancy_rate_raw_text"].tolist())
                            if value != ""
                        )
                    ),
                    "source_occurrences": int(len(group)),
                    "merged_source_rows": "|".join(str(int(value)) for value in group["source_row"].tolist()),
                    "rooms_sold": int(kept["rooms_sold"]),
                    "transient_rooms": int(kept["transient_rooms"]),
                    "group_rooms": int(kept["group_rooms"]),
                    "day_use_rooms": int(kept["day_use_rooms"]),
                    "comp_rooms": int(kept["comp_rooms"]),
                    "out_of_order_rooms": int(kept["out_of_order_rooms"]),
                    "on_market_rooms": int(kept["on_market_rooms"]),
                    "room_revenue_gross": kept["room_revenue_gross"],
                    "adr_gross": kept["adr_gross"],
                    "revpar_gross": kept["revpar_gross"],
                    "room_revenue_net": kept["room_revenue_net"],
                    "adr_net": kept["adr_net"],
                    "revpar_net": kept["revpar_net"],
                }
            )
            for _, dropped in group.iloc[1:].iterrows():
                duplicate_rows.append(
                    {
                        "dataset_name": "block_roomsales_adjustments",
                        "source_file": dropped["source_file"],
                        "source_sheet": dropped["source_sheet"],
                        "source_row": int(dropped["source_row"]),
                        "business_date": dropped["business_date"],
                        "dedupe_key_columns": "|".join(adjustment_group_cols),
                        "dedupe_key": "|".join(str(dropped[column]) for column in adjustment_group_cols),
                        "kept_source_file": kept["source_file"],
                        "kept_source_sheet": kept["source_sheet"],
                        "kept_source_row": int(kept["source_row"]),
                        "drop_reason": "duplicate special adjustment row with same date and revenue metrics; duplicate pair retained once",
                    }
                )

    adjustment_df = pd.DataFrame(adjustment_clean_records).sort_values(
        ["business_date", "source_file", "source_row"]
    ).reset_index(drop=True)
    duplicate_df = pd.DataFrame(
        duplicate_rows,
        columns=[
            "dataset_name",
            "source_file",
            "source_sheet",
            "source_row",
            "business_date",
            "dedupe_key_columns",
            "dedupe_key",
            "kept_source_file",
            "kept_source_sheet",
            "kept_source_row",
            "drop_reason",
        ],
    )

    daily_df.to_csv(BLOCK_DAILY_OUTPUT, index=False)
    class_df.to_csv(BLOCK_CLASS_OUTPUT, index=False)
    type_df.to_csv(BLOCK_TYPE_OUTPUT, index=False)
    adjustment_df.to_csv(BLOCK_ADJUSTMENT_OUTPUT, index=False)
    duplicate_df.to_csv(BLOCK_DUPES, index=False)

    class_rollup_checks: dict[str, dict[str, int]] = {}
    for room_class, labels in {
        "Standard": ["HKNS", "KFNS", "KNS", "QQNS"],
        "Suite": ["DKSNS", "HKOBNS", "KESNS", "KOBNS", "KSSNS", "QQOBNS"],
    }.items():
        class_subset = class_df[class_df["room_class"] == room_class].sort_values("business_date")
        type_rollup = (
            type_df[type_df["room_type_code"].isin(labels)]
            .groupby("business_date", dropna=False)[
                [
                    "rooms_sold",
                    "transient_rooms",
                    "group_rooms",
                    "day_use_rooms",
                    "comp_rooms",
                    "out_of_order_rooms",
                    "on_market_rooms",
                    "room_revenue_gross",
                    "room_revenue_net",
                ]
            ]
            .sum(min_count=1)
            .reset_index()
        )
        merged = class_subset.merge(type_rollup, on="business_date", suffixes=("_class", "_types"))
        mismatch_counts: dict[str, int] = {}
        for column in [
            "rooms_sold",
            "transient_rooms",
            "group_rooms",
            "day_use_rooms",
            "comp_rooms",
            "out_of_order_rooms",
            "on_market_rooms",
        ]:
            mismatch_counts[column] = int((merged[f"{column}_class"] != merged[f"{column}_types"]).sum())
        for column in ["room_revenue_gross", "room_revenue_net"]:
            mismatch_counts[column] = int(
                (merged[f"{column}_class"].round(2) != merged[f"{column}_types"].round(2)).sum()
            )
        class_rollup_checks[room_class] = mismatch_counts

    combined_components = pd.concat(
        [
            type_df[
                [
                    "business_date",
                    "rooms_sold",
                    "transient_rooms",
                    "group_rooms",
                    "day_use_rooms",
                    "comp_rooms",
                    "out_of_order_rooms",
                    "on_market_rooms",
                    "room_revenue_gross",
                    "room_revenue_net",
                ]
            ],
            adjustment_df[
                [
                    "business_date",
                    "rooms_sold",
                    "transient_rooms",
                    "group_rooms",
                    "day_use_rooms",
                    "comp_rooms",
                    "out_of_order_rooms",
                    "on_market_rooms",
                    "room_revenue_gross",
                    "room_revenue_net",
                ]
            ],
        ],
        ignore_index=True,
    )
    component_rollup = (
        combined_components.groupby("business_date", dropna=False)[
            [
                "rooms_sold",
                "transient_rooms",
                "group_rooms",
                "day_use_rooms",
                "comp_rooms",
                "out_of_order_rooms",
                "on_market_rooms",
                "room_revenue_gross",
                "room_revenue_net",
            ]
        ]
        .sum(min_count=1)
        .reset_index()
    )
    reconciliation = daily_df.merge(component_rollup, on="business_date", suffixes=("_daily", "_components"))
    reconciliation_counts: dict[str, int] = {}
    for column in [
        "rooms_sold",
        "transient_rooms",
        "group_rooms",
        "day_use_rooms",
        "comp_rooms",
        "out_of_order_rooms",
        "on_market_rooms",
    ]:
        reconciliation_counts[column] = int(
            (reconciliation[f"{column}_daily"] != reconciliation[f"{column}_components"]).sum()
        )
    for column in ["room_revenue_gross", "room_revenue_net"]:
        reconciliation_counts[column] = int(
            (
                reconciliation[f"{column}_daily"].round(2)
                != reconciliation[f"{column}_components"].round(2)
            ).sum()
        )

    summary = {
        "dataset_family": "block_roomsales",
        "raw_files_used": [meta["source_file"] for meta in file_meta],
        "cleaned_files_created": [
            BLOCK_DAILY_OUTPUT.name,
            BLOCK_CLASS_OUTPUT.name,
            BLOCK_TYPE_OUTPUT.name,
            BLOCK_ADJUSTMENT_OUTPUT.name,
            BLOCK_DUPES.name,
            BLOCK_SUMMARY.name,
        ],
        "report_layout": {
            "sheet_name": "RoomSales",
            "header_rows_before_daily_data": 10,
            "daily_total_row_rule": "Rows beginning at row 11 where column A is a valid Excel date were treated as property-level daily totals.",
            "room_class_row_rule": "Rows with labels Standard or Suite immediately under a business_date were treated as room-class rollups.",
            "room_type_row_rule": "Rows with room-type codes HKNS, KFNS, KNS, QQNS, DKSNS, HKOBNS, KESNS, KOBNS, KSSNS, and QQOBNS were treated as leaf room-type rows.",
            "adjustment_row_rule": "Rows labeled *** were isolated as special adjustment rows because they are not valid room types and can carry revenue-only adjustments.",
            "summary_section_rule": "Once a row starting with Summary - was reached, that row and all following room-class totals were excluded from the cleaned daily tables.",
            "footer_rule": "Rows beginning with *Totals for past dates were excluded as report footnotes.",
        },
        "date_coverage": {
            "start": daily_df["business_date"].min(),
            "end": daily_df["business_date"].max(),
            "daily_rows_cleaned": int(len(daily_df)),
            "room_class_rows_cleaned": int(len(class_df)),
            "room_type_rows_cleaned": int(len(type_df)),
            "adjustment_rows_cleaned": int(len(adjustment_df)),
            "missing_business_dates_after_cleaning": missing_daily_dates(daily_df["business_date"].tolist()),
            "per_file": file_meta,
            "property_total_rooms_inferred_unique_values": sorted(
                {int(meta["property_total_rooms_inferred"]) for meta in file_meta}
            ),
        },
        "transformations": [
            "Converted business dates from Excel date serial / datetime storage to ISO YYYY-MM-DD.",
            "Inferred daily property_total_rooms by dividing the Total Rooms period total shown in row 10 by the number of report days; every file resolved to 60 rooms.",
            "Normalized Occ % to decimal occupancy_rate_reported rounded to 4 decimals.",
            "Normalized gross and net room revenue to numeric values rounded to 2 decimals and ADR / RevPAR values to numeric values rounded to 4 decimals.",
            "Converted room-count fields to nullable integers.",
            "Split the mixed RoomSales report into separate cleaned outputs for property totals, room-class rollups, leaf room-type rows, and special adjustment rows so aggregate class rows are not double-counted as room-type rows.",
            "Mapped room-type codes to their parent room_class based on the report structure: Standard -> HKNS/KFNS/KNS/QQNS and Suite -> DKSNS/HKOBNS/KESNS/KOBNS/KSSNS/QQOBNS.",
            "Preserved special *** rows in a dedicated adjustments dataset and set occupancy_rate_reported to null there because the raw occupancy values are not meaningful inventory metrics for those rows.",
        ],
        "deduplication": {
            "daily_totals_key": ["business_date"],
            "room_class_key": ["business_date", "room_class"],
            "room_type_key": ["business_date", "room_type_code"],
            "adjustment_key": adjustment_group_cols,
            "duplicate_rows_removed_total": int(len(duplicate_df)),
            "duplicate_rows_removed_by_dataset": {
                "block_roomsales_daily_totals": int((duplicate_df["dataset_name"] == "block_roomsales_daily_totals").sum())
                if not duplicate_df.empty
                else 0,
                "block_roomsales_room_class_daily": int(
                    (duplicate_df["dataset_name"] == "block_roomsales_room_class_daily").sum()
                )
                if not duplicate_df.empty
                else 0,
                "block_roomsales_room_type_daily": int(
                    (duplicate_df["dataset_name"] == "block_roomsales_room_type_daily").sum()
                )
                if not duplicate_df.empty
                else 0,
                "block_roomsales_adjustments": int(
                    (duplicate_df["dataset_name"] == "block_roomsales_adjustments").sum()
                )
                if not duplicate_df.empty
                else 0,
            },
            "note": "The only duplicate rows found in the Block family were repeated pairs of special *** adjustment rows. One copy of each pair was retained and the removed copies were logged.",
        },
        "rows_removed_or_excluded": {
            "summary_section_rows_excluded": summary_section_rows_removed,
            "footnote_rows_excluded": footnote_rows_removed,
            "raw_adjustment_rows_detected": int(len(adjustment_raw_df)),
            "adjustment_rows_after_deduplication": int(len(adjustment_df)),
        },
        "validation": {
            "room_class_rollup_mismatch_counts": class_rollup_checks,
            "property_daily_reconciliation_mismatch_counts_using_room_type_rows_plus_deduped_adjustments": reconciliation_counts,
            "room_type_rows_per_day_min": int(type_df.groupby("business_date")["room_type_code"].nunique().min()),
            "room_type_rows_per_day_max": int(type_df.groupby("business_date")["room_type_code"].nunique().max()),
        },
        "ambiguous_or_inconsistent_filenames": [
            "Block 4 October 4 to December 31.xls does not contain a year in the filename; the internal report range confirms it covers 2023-10-04 through 2023-12-31.",
            "Block12023roomtypes.xls, Block22023roomtypes.xls, and Block32023roomtypes.xls do not spell out their date ranges in the filename; the internal report header provides the exact coverage.",
            "Block 1  92 daysJanuary 1 to April 1 2024.xls contains irregular spacing and wording in the filename; the internal report header confirms the exact range 2024-01-01 through 2024-04-01.",
        ],
        "data_quality_notes": [
            "The Block RoomSales Occ % metric is based on total property inventory (60 rooms), not on-market rooms.",
            "Special *** adjustment rows occur on 10 unique business dates across 2023-2025 and can carry revenue-only adjustments. They were preserved separately so room-class and room-type outputs remain analytically safe.",
        ],
        "new_data_capabilities_from_this_family": [
            "Continuous daily property-level operational history from 2023-01-01 through 2025-12-31.",
            "Explicit daily on-market rooms, out-of-order rooms, comp rooms, day-use rooms, transient rooms, and group rooms.",
            "Daily gross and net room revenue plus gross and net ADR / RevPAR.",
            "Daily room-class rollups for Standard and Suite.",
            "Daily leaf room-type metrics for 10 room-type codes with parent room-class mapping.",
        ],
    }
    BLOCK_SUMMARY.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return daily_df, class_df, type_df, adjustment_df, duplicate_df, summary


def add_cross_family_validation(
    reservation_clean: pd.DataFrame,
    block_daily: pd.DataFrame,
    block_summary: dict,
) -> None:
    shared = reservation_clean.merge(block_daily, on="business_date", how="inner", suffixes=("_forecast", "_block"))
    shared["rooms_sold_over_on_market"] = (
        shared["rooms_sold_forecast"] / shared["on_market_rooms"]
    ).round(2)
    block_summary["validation"]["cross_family_comparison_with_reservation_forecast_2023_2024"] = {
        "shared_business_dates": int(len(shared)),
        "exact_match_counts": {
            "rooms_sold": int((shared["rooms_sold_forecast"] == shared["rooms_sold_block"]).sum()),
            "out_of_order_rooms": int(
                (shared["out_of_order_rooms_forecast"] == shared["out_of_order_rooms_block"]).sum()
            ),
            "room_revenue_net": int(
                (shared["room_revenue_net_forecast"].round(2) == shared["room_revenue_net_block"].round(2)).sum()
            ),
        },
        "match_counts_after_rounding_to_2_decimals": {
            "adr_reported_vs_adr_net": int(
                (shared["adr_reported"].round(2) == shared["adr_net"].round(2)).sum()
            ),
            "revpar_reported_vs_revpar_net": int(
                (shared["revpar_reported"].round(2) == shared["revpar_net"].round(2)).sum()
            ),
        },
        "notable_differences": {
            "transient_room_count_differences": int(
                (shared["transient_rooms_forecast"] != shared["transient_rooms_block"]).sum()
            ),
            "group_room_count_differences": int(
                (shared["group_rooms_picked_up"] != shared["group_rooms"]).sum()
            ),
        },
        "occupancy_basis_inference": {
            "forecast_occ_matches_rooms_sold_divided_by_on_market_when_rounded_to_2_decimals": int(
                (
                    shared["occupancy_rate_reported_forecast"].round(2)
                    == shared["rooms_sold_over_on_market"]
                ).sum()
            ),
            "total_shared_dates": int(len(shared)),
            "note": "This suggests ReservationForecast Occ% is effectively based on on-market rooms, while Block RoomSales Occ % is based on total property rooms.",
        },
    }
    block_summary["data_quality_notes"].append(
        "ReservationForecast Occ% closely matches rooms_sold / on_market on overlapping dates when rounded to 2 decimals, so the two families appear to use different occupancy denominators."
    )
    BLOCK_SUMMARY.write_text(json.dumps(block_summary, indent=2), encoding="utf-8")


def main() -> None:
    CLEANED_DIR.mkdir(exist_ok=True)
    staging_map = ensure_staging_files(raw_xls_files())
    reservation_clean, reservation_dupes, reservation_summary = parse_reservation_forecast(staging_map)
    block_daily, block_class, block_type, block_adjustments, block_dupes, block_summary = parse_block_roomsales(
        staging_map
    )
    add_cross_family_validation(reservation_clean, block_daily, block_summary)

    print(
        json.dumps(
            {
                "created_files": [
                    RESERVATION_OUTPUT.name,
                    RESERVATION_DUPES.name,
                    RESERVATION_SUMMARY.name,
                    BLOCK_DAILY_OUTPUT.name,
                    BLOCK_CLASS_OUTPUT.name,
                    BLOCK_TYPE_OUTPUT.name,
                    BLOCK_ADJUSTMENT_OUTPUT.name,
                    BLOCK_DUPES.name,
                    BLOCK_SUMMARY.name,
                ],
                "reservation_forecast_rows": int(len(reservation_clean)),
                "reservation_forecast_duplicates_removed": int(len(reservation_dupes)),
                "block_daily_rows": int(len(block_daily)),
                "block_room_class_rows": int(len(block_class)),
                "block_room_type_rows": int(len(block_type)),
                "block_adjustment_rows": int(len(block_adjustments)),
                "block_duplicates_removed": int(len(block_dupes)),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
